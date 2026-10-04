"""
Update S3E05 cards with Chinese translations from Excel vocabulary sheets.
Includes both word translations and sentence translations.
"""

import sys
import io
import json
from pathlib import Path
from collections import defaultdict

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

import openpyxl

PROJECT_ROOT = Path(__file__).parent.parent
EXCEL_PATH = PROJECT_ROOT / "Money Heist" / "La_Casa_de_Papel_S3E01_台词.xlsx"
CARDS_PATH = PROJECT_ROOT / "src" / "data" / "LCDP_S3E05_cards.json"


def load_vocab_from_excel():
    """Load all S3E05 vocabulary from Excel sheets with translations."""
    wb = openpyxl.load_workbook(str(EXCEL_PATH), data_only=True)
    
    vocab = {}  # word_lower -> {translation, sentence_es, sentence_zh}
    
    # Sheets to process for S3E05
    sheets_to_process = [
        'S3E05 动词', 'S3E05 形容词', 'S3E05 副词', 'S3E05 名词',
        'S3E05E07 动词', 'S3E05E07 名词', 'S3E05E07 副词', 'S3E05E07 形容词'
    ]
    
    for sheet_name in sheets_to_process:
        if sheet_name not in wb.sheetnames:
            print(f"Sheet not found: {sheet_name}", flush=True)
            continue
            
        ws = wb[sheet_name]
        count = 0
        
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or not row[0]:
                continue
                
            # Different sheets have different column layouts
            if '动词' in sheet_name:
                # 序号, 动词原形, 中文释义, 剧中变位形式, 例句(西), 例句(中)
                if len(row) >= 6:
                    lemma = str(row[1]).strip().lower() if row[1] else ''
                    translation = str(row[2]).strip() if row[2] else ''
                    forms = str(row[3]).strip() if row[3] else ''
                    sentence_es = str(row[4]).strip() if row[4] else ''
                    sentence_zh = str(row[5]).strip() if row[5] else ''
                    
                    # Store lemma
                    if lemma and translation:
                        vocab[lemma] = {
                            'translation': translation,
                            'sentence_es': sentence_es,
                            'sentence_zh': sentence_zh
                        }
                        count += 1
                    
                    # Also store inflected forms if present
                    if forms:
                        for form in forms.split(','):
                            form = form.strip().lower()
                            if form and form != lemma:
                                vocab[form] = {
                                    'translation': translation,
                                    'sentence_es': sentence_es,
                                    'sentence_zh': sentence_zh
                                }
            
            elif '名词' in sheet_name:
                # 序号, 名词, 中文释义, 出现次数, 例句(西), 例句(中)
                if len(row) >= 6:
                    word = str(row[1]).strip().lower() if row[1] else ''
                    translation = str(row[2]).strip() if row[2] else ''
                    sentence_es = str(row[4]).strip() if row[4] else ''
                    sentence_zh = str(row[5]).strip() if row[5] else ''
                    
                    if word and translation:
                        vocab[word] = {
                            'translation': translation,
                            'sentence_es': sentence_es,
                            'sentence_zh': sentence_zh
                        }
                        count += 1
            
            elif '形容词' in sheet_name:
                # 序号, 原形, 中文释义, 剧中形式, 例句(西), 例句(中)
                if len(row) >= 6:
                    base = str(row[1]).strip().lower() if row[1] else ''
                    translation = str(row[2]).strip() if row[2] else ''
                    forms = str(row[3]).strip() if row[3] else ''
                    sentence_es = str(row[4]).strip() if row[4] else ''
                    sentence_zh = str(row[5]).strip() if row[5] else ''
                    
                    if base and translation:
                        vocab[base] = {
                            'translation': translation,
                            'sentence_es': sentence_es,
                            'sentence_zh': sentence_zh
                        }
                        count += 1
                    
                    # Also store inflected forms
                    if forms:
                        for form in forms.split(','):
                            form = form.strip().lower()
                            if form and form != base:
                                vocab[form] = {
                                    'translation': translation,
                                    'sentence_es': sentence_es,
                                    'sentence_zh': sentence_zh
                                }
            
            elif '副词' in sheet_name:
                # 序号, 副词, 中文释义, 剧中形式, 例句(西), 例句(中)
                if len(row) >= 6:
                    base = str(row[1]).strip().lower() if row[1] else ''
                    translation = str(row[2]).strip() if row[2] else ''
                    forms = str(row[3]).strip() if row[3] else ''
                    sentence_es = str(row[4]).strip() if row[4] else ''
                    sentence_zh = str(row[5]).strip() if row[5] else ''
                    
                    if base and translation:
                        vocab[base] = {
                            'translation': translation,
                            'sentence_es': sentence_es,
                            'sentence_zh': sentence_zh
                        }
                        count += 1
                    
                    # Also store inflected forms
                    if forms:
                        for form in forms.split(','):
                            form = form.strip().lower()
                            if form and form != base:
                                vocab[form] = {
                                    'translation': translation,
                                    'sentence_es': sentence_es,
                                    'sentence_zh': sentence_zh
                                }
        
        print(f"Loaded {count} entries from {sheet_name}", flush=True)
    
    return vocab


def load_dialogue_translations():
    """Load sentence translations from dialogue sheet."""
    wb = openpyxl.load_workbook(str(EXCEL_PATH), data_only=True)
    
    dialogue_trans = {}  # spanish_sentence -> chinese_translation
    
    # Find the dialogue sheet for S3E05
    dialogue_sheet = None
    for name in wb.sheetnames:
        if '纸钞屋' in name and 'S3E05' in name and '台词' in name:
            dialogue_sheet = name
            break
    
    if not dialogue_sheet:
        print("Dialogue sheet not found for S3E05", flush=True)
        return dialogue_trans
    
    ws = wb[dialogue_sheet]
    count = 0
    
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row and len(row) >= 3:
            spanish = str(row[1]).strip() if row[1] else ''
            chinese = str(row[2]).strip() if row[2] else ''
            
            if spanish and chinese:
                dialogue_trans[spanish] = chinese
                count += 1
    
    print(f"Loaded {count} dialogue translations", flush=True)
    return dialogue_trans


def update_cards(vocab, dialogue_trans):
    """Update S3E05 cards with translations."""
    with open(CARDS_PATH, 'r', encoding='utf-8') as f:
        cards = json.load(f)
    
    updated_word_trans = 0
    updated_sentence_trans = 0
    matched_dialogue = 0
    
    for card in cards:
        target_word = card['target_word'].lower()
        sentence_full = card.get('sentence_full', '').strip()
        
        # Update word translation
        if target_word in vocab:
            entry = vocab[target_word]
            if entry['translation'] and not card.get('translation', '').strip():
                card['translation'] = entry['translation']
                updated_word_trans += 1
            
            # Update sentence translation from vocab
            if entry['sentence_zh'] and not card.get('sentence_translation', '').strip():
                card['sentence_translation'] = entry['sentence_zh']
                updated_sentence_trans += 1
        
        # If still no sentence translation, try matching with dialogue
        if not card.get('sentence_translation', '').strip() and sentence_full:
            if sentence_full in dialogue_trans:
                card['sentence_translation'] = dialogue_trans[sentence_full]
                matched_dialogue += 1
    
    # Save updated cards
    with open(CARDS_PATH, 'w', encoding='utf-8') as f:
        json.dump(cards, f, ensure_ascii=False, indent=2)
    
    print(f"\nUpdated {updated_word_trans} word translations", flush=True)
    print(f"Updated {updated_sentence_trans} sentence translations from vocab", flush=True)
    print(f"Matched {matched_dialogue} sentence translations from dialogue", flush=True)
    
    # Final stats
    with open(CARDS_PATH, 'r', encoding='utf-8') as f:
        cards = json.load(f)
    
    has_word_trans = sum(1 for c in cards if c.get('translation', '').strip())
    has_sent_trans = sum(1 for c in cards if c.get('sentence_translation', '').strip())
    
    print(f"\nFinal stats:", flush=True)
    print(f"  Word translations: {has_word_trans}/{len(cards)}", flush=True)
    print(f"  Sentence translations: {has_sent_trans}/{len(cards)}", flush=True)


if __name__ == '__main__':
    print("Loading vocabulary from Excel...", flush=True)
    vocab = load_vocab_from_excel()
    print(f"\nTotal vocabulary entries: {len(vocab)}", flush=True)
    
    print("\nLoading dialogue translations...", flush=True)
    dialogue_trans = load_dialogue_translations()
    
    print("\nUpdating cards...", flush=True)
    update_cards(vocab, dialogue_trans)
    
    print("\nDone!", flush=True)
