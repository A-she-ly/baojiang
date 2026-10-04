import sys
import io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

import openpyxl

wb = openpyxl.load_workbook('Money Heist/La_Casa_de_Papel_S3E01_台词.xlsx', data_only=True)

# Check S3E05E07 形容词 sheet
ws = wb['S3E05E07 形容词']
print("Header:", [cell.value for cell in ws[1]])
for i, row in enumerate(ws.iter_rows(min_row=2, max_row=5, values_only=True), 2):
    print(f"Row {i}: {row}")
