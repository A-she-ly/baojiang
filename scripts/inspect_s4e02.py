import openpyxl
import sys
import io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

wb = openpyxl.load_workbook('Money Heist/La_Casa_de_Papel_S3E01_台词.xlsx', read_only=True)

# Find S4E02 sheets
s4e02_sheets = [s for s in wb.sheetnames if 'S4E02' in s or 'S4E02' in s]
print(f"S4E02 sheets: {s4e02_sheets}")

for sheet_name in s4e02_sheets:
    ws = wb[sheet_name]
    rows = list(ws.iter_rows(max_row=5, values_only=True))
    print(f"\n=== {sheet_name} ===")
    print(f"Rows: {ws.max_row}, Cols: {ws.max_column}")
    for i, row in enumerate(rows):
        print(f"  Row {i}: {row}")

wb.close()
