"""
Fix S3E05 cloze fields using Excel vocabulary data.
The Excel has a column with inflected forms that appear in the show.
"""

import sys
import io
import json
import re
import unicodedata
from pathlib import Path

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

import openpyxl

PROJECT_ROOT = Path(__file__).parent.parent
EXCEL_PATH = PROJECT_ROOT / "Money Heist" / "La_Casa_de_Papel_S3E01_台词.xlsx"
CARDS_PATH = PROJECT_ROOT / "src" / "data" / "LCDP_S3E05_cards.json"


def remove_accents(text):
    """Remove accents from text for comparison."""
    nfkd = unicodedata.normalize('NFKD', text)
    return ''.join(c for c in nfkd if not unicodedata.combining(c))


def load_excel_forms():
    """Load inflected forms from Excel vocabulary sheets."""
    wb = openpyxl.load_workbook(str(EXCEL_PATH), data_only=True)
    
    # Map: lemma_lower -> set of inflected forms
    forms_map = {}
    
    sheets_to_process = [
        'S3E05 动词', 'S3E05 形容词', 'S3E05 副词', 'S3E05 名词',
        'S3E05E07 动词', 'S3E05E07 名词', 'S3E05E07 副词', 'S3E05E07 形容词'
    ]
    
    for sheet_name in sheets_to_process:
        if sheet_name not in wb.sheetnames:
            continue
            
        ws = wb[sheet_name]
        
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or not row[0]:
                continue
            
            if '动词' in sheet_name:
                # 序号, 动词原形, 中文释义, 剧中变位形式, 例句(西), 例句(中)
                if len(row) >= 4 and row[1] and row[3]:
                    lemma = str(row[1]).strip().lower()
                    forms_str = str(row[3]).strip()
                    forms = [f.strip().lower() for f in forms_str.split(',')]
                    forms_map[lemma] = forms
            
            elif '形容词' in sheet_name or '副词' in sheet_name:
                # 序号, 原形, 中文释义, 剧中形式, 例句(西), 例句(中)
                if len(row) >= 4 and row[1] and row[3]:
                    base = str(row[1]).strip().lower()
                    forms_str = str(row[3]).strip()
                    forms = [f.strip().lower() for f in forms_str.split(',')]
                    forms_map[base] = forms
            
            elif '名词' in sheet_name:
                # 序号, 名词, 中文释义, 出现次数, 例句(西), 例句(中)
                # Nouns don't have a forms column, use the word itself
                if len(row) >= 2 and row[1]:
                    word = str(row[1]).strip().lower()
                    forms_map[word] = [word]
    
    return forms_map


def find_word_in_sentence(sentence, target_word, forms_map):
    """Find the actual word form in the sentence."""
    target_lower = target_word.lower()
    
    # Get possible forms for this target
    possible_forms = [target_lower]
    if target_lower in forms_map:
        possible_forms.extend(forms_map[target_lower])
    
    # Also add accent-insensitive versions
    target_no_accents = remove_accents(target_lower)
    
    # Find all words in sentence
    words = re.findall(r'\b\w+\b', sentence)
    
    for word in words:
        word_lower = word.lower()
        word_no_accents = remove_accents(word_lower)
        
        # Check if this word matches any possible form
        if word_lower in possible_forms:
            return word
        
        # Check accent-insensitive match
        if word_no_accents == target_no_accents:
            return word
        
        # Check if word matches any inflected form (accent-insensitive)
        for form in possible_forms:
            if remove_accents(form) == word_no_accents:
                return word
    
    return None


def fix_clozes():
    """Fix all cloze fields in S3E05 cards."""
    print("Loading forms from Excel...", flush=True)
    forms_map = load_excel_forms()
    print(f"Loaded forms for {len(forms_map)} lemmas", flush=True)
    
    with open(CARDS_PATH, 'r', encoding='utf-8') as f:
        cards = json.load(f)
    
    fixed_count = 0
    not_found = []
    
    for card in cards:
        target_word = card.get('target_word', '')
        sentence_full = card.get('sentence_full', '')
        sentence_cloze = card.get('sentence_cloze', '')
        
        # Check if cloze needs fixing
        if '___' not in sentence_cloze and target_word and sentence_full:
            # Find the actual word form in the sentence
            actual_word = find_word_in_sentence(sentence_full, target_word, forms_map)
            
            if actual_word:
                # Replace the word with ___
                new_cloze = sentence_full.replace(actual_word, '___', 1)
                if new_cloze != sentence_cloze:
                    card['sentence_cloze'] = new_cloze
                    fixed_count += 1
            else:
                not_found.append((target_word, sentence_full[:50]))
    
    # Save updated cards
    with open(CARDS_PATH, 'w', encoding='utf-8') as f:
        json.dump(cards, f, ensure_ascii=False, indent=2)
    
    print(f"\nFixed {fixed_count} cloze fields", flush=True)
    print(f"Could not find word for {len(not_found)} cards:", flush=True)
    for target, sentence in not_found[:20]:
        print(f"  {target}: {sentence}...", flush=True)


if __name__ == '__main__':
    fix_clozes()
