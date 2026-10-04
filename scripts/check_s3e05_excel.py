import sys
import io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

import openpyxl

wb = openpyxl.load_workbook('Money Heist/La_Casa_de_Papel_S3E01_台词.xlsx', data_only=True)

# Check all S3E05 related sheets
for name in wb.sheetnames:
    if 'S3E05' in name:
        ws = wb[name]
        print(f"\n=== {name} ===")
        print(f"Rows: {ws.max_row}, Cols: {ws.max_column}")
        # Print header
        header = [cell.value for cell in ws[1]]
        print(f"Header: {header}")
        # Print first 3 data rows
        for i, row in enumerate(ws.iter_rows(min_row=2, max_row=4, values_only=True), 2):
            print(f"Row {i}: {row}")
