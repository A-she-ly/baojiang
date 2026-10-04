"""
Fix S3E07 target_word fields to include proper accents.
The Excel vocabulary sheets have the correct accented forms.
"""

import sys
import io
import json
import unicodedata
from pathlib import Path

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

import openpyxl

PROJECT_ROOT = Path(__file__).parent.parent
EXCEL_PATH = PROJECT_ROOT / "Money Heist" / "La_Casa_de_Papel_S3E01_台词.xlsx"
CARDS_PATH = PROJECT_ROOT / "src" / "data" / "LCDP_S3E07_cards.json"


def remove_accents(text):
    """Remove accents from text for comparison."""
    nfkd = unicodedata.normalize('NFKD', text)
    return ''.join(c for c in nfkd if not unicodedata.combining(c))


def load_correct_forms():
    """Load correct accented forms from Excel vocabulary sheets."""
    wb = openpyxl.load_workbook(str(EXCEL_PATH), data_only=True)
    
    # Map: unaccented_lower -> accented_form
    correct_forms = {}
    
    sheets_to_process = [
        'S3E05E07 动词', 'S3E05E07 名词', 'S3E05E07 副词', 'S3E05E07 形容词'
    ]
    
    for sheet_name in sheets_to_process:
        if sheet_name not in wb.sheetnames:
            continue
            
        ws = wb[sheet_name]
        
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or not row[0]:
                continue
            
            if '动词' in sheet_name:
                # 序号, 动词原形, 中文释义, 剧中变位形式, 例句(西), 例句(中)
                if len(row) >= 4 and row[1]:
                    lemma = str(row[1]).strip()
                    forms_str = str(row[3]).strip() if row[3] else ''
                    
                    # Store lemma
                    unaccented = remove_accents(lemma.lower())
                    correct_forms[unaccented] = lemma.lower()
                    
                    # Store forms
                    if forms_str:
                        for form in forms_str.split(','):
                            form = form.strip()
                            unaccented_form = remove_accents(form.lower())
                            correct_forms[unaccented_form] = form.lower()
            
            elif '名词' in sheet_name:
                # 序号, 名词, 中文释义, 出现次数, 例句(西), 例句(中)
                if len(row) >= 2 and row[1]:
                    word = str(row[1]).strip()
                    unaccented = remove_accents(word.lower())
                    correct_forms[unaccented] = word.lower()
            
            elif '形容词' in sheet_name or '副词' in sheet_name:
                # 序号, 原形, 中文释义, 剧中形式, 例句(西), 例句(中)
                if len(row) >= 4 and row[1]:
                    base = str(row[1]).strip()
                    forms_str = str(row[3]).strip() if row[3] else ''
                    
                    # Store base
                    unaccented = remove_accents(base.lower())
                    correct_forms[unaccented] = base.lower()
                    
                    # Store forms
                    if forms_str:
                        for form in forms_str.split(','):
                            form = form.strip()
                            unaccented_form = remove_accents(form.lower())
                            correct_forms[unaccented_form] = form.lower()
    
    return correct_forms


def fix_target_words():
    """Fix target_word fields in S3E07 cards."""
    print("Loading correct forms from Excel...", flush=True)
    correct_forms = load_correct_forms()
    print(f"Loaded {len(correct_forms)} correct forms", flush=True)
    
    with open(CARDS_PATH, 'r', encoding='utf-8') as f:
        cards = json.load(f)
    
    fixed_count = 0
    for card in cards:
        target_word = card.get('target_word', '')
        unaccented = remove_accents(target_word.lower())
        
        # Check if we have a correct accented form
        if unaccented in correct_forms:
            correct_form = correct_forms[unaccented]
            if correct_form != target_word.lower():
                card['target_word'] = correct_form
                # Also update pos field
                pos_prefix = card['pos'].split()[0] if ' ' in card['pos'] else 'n.'
                card['pos'] = f"{pos_prefix} {correct_form}"
                fixed_count += 1
    
    # Save updated cards
    with open(CARDS_PATH, 'w', encoding='utf-8') as f:
        json.dump(cards, f, ensure_ascii=False, indent=2)
    
    print(f"\nFixed {fixed_count} target_word fields", flush=True)


if __name__ == '__main__':
    fix_target_words()
