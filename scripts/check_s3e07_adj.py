import sys
import io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

import openpyxl

wb = openpyxl.load_workbook('Money Heist/La_Casa_de_Papel_S3E01_台词.xlsx', data_only=True)
print("Sheet names:", wb.sheetnames)

# Find S3E07 adjective sheet
for name in wb.sheetnames:
    if 'S3E05E07' in name and '形容' in name:
        ws = wb[name]
        print(f"\n=== {name} ===")
        print(f"Rows: {ws.max_row}, Cols: {ws.max_column}")
        for row in ws.iter_rows(min_row=1, max_row=10, values_only=True):
            print(row)
